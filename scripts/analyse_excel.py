"""Profile the MarketingCollateral workbook (all sheets, including hidden).

Read-only. Prints a Markdown report to stdout. Redirect it to docs/ or the
scratchpad. Usage:

    python -I scripts/analyse_excel.py data/MarketingCollateral.xlsx > report.md
    python -I scripts/analyse_excel.py data/MarketingCollateral.xlsx --dump

--dump prints every non-empty row of every sheet (use it to review raw data).
"""

import argparse
import datetime as dt
import re
import sys
from collections import Counter

from openpyxl import load_workbook

EXCEL_EPOCH = dt.datetime(1899, 12, 30)
SERIAL_RANGE = (30000, 60000)  # ~1982..2064; anything outside is not a date serial


def clean(value):
    if isinstance(value, str):
        return value.strip()
    return value


def is_serial_date(value):
    return isinstance(value, (int, float)) and SERIAL_RANGE[0] <= value <= SERIAL_RANGE[1]


def serial_to_iso(value):
    return (EXCEL_EPOCH + dt.timedelta(days=float(value))).date().isoformat()


def header_row(rows):
    """First row with at least two non-empty cells is treated as the header."""
    for idx, row in enumerate(rows):
        filled = [c for c in row if clean(c) not in (None, "")]
        if len(filled) >= 2:
            return idx
    return None


def sheet_rows(ws):
    return [tuple(clean(c.value) for c in row) for row in ws.iter_rows()]


def profile_sheet(ws, out):
    rows = sheet_rows(ws)
    non_empty = [(i + 1, r) for i, r in enumerate(rows) if any(c not in (None, "") for c in r)]
    out.append(f"### `{ws.title}`")
    out.append("")
    out.append(f"- State: **{ws.sheet_state}**")
    out.append(f"- Dimensions: {ws.dimensions}; non-empty rows: {len(non_empty)}")
    if not non_empty:
        out.append("- Empty sheet.")
        out.append("")
        return

    hidx = header_row(rows)
    if hidx is None:
        out.append("- No header row detected.")
        out.append("")
        return
    header = [str(h) if h not in (None, "") else f"<col {i + 1}>" for i, h in enumerate(rows[hidx])]
    out.append(f"- Header (row {hidx + 1}): " + " | ".join(f"`{h}`" for h in header))
    out.append("")

    data = [(i + 1, r) for i, r in enumerate(rows) if i > hidx and any(c not in (None, "") for c in r)]
    out.append(f"- Data rows after header: {len(data)}")
    out.append("")
    out.append("| Column | Non-empty | Types seen | Date-serial values | Distinct values (top 6) |")
    out.append("|---|---|---|---|---|")
    for col_idx, name in enumerate(header):
        vals = [r[col_idx] if col_idx < len(r) else None for _, r in data]
        present = [v for v in vals if v not in (None, "")]
        types = Counter(type(v).__name__ for v in present)
        serials = sum(1 for v in present if is_serial_date(v))
        distinct = Counter(str(v) for v in present).most_common(6)
        dist_txt = ", ".join(f"`{d[:40]}` ×{n}" for d, n in distinct)
        out.append(
            f"| {name} | {len(present)}/{len(vals)} | {dict(types)} | {serials} | {dist_txt} |"
        )
    out.append("")


def find_issues(ws, out):
    """Cheap, rule-based checks. Findings are candidates to confirm, not fixes."""
    rows = sheet_rows(ws)
    hidx = header_row(rows)
    if hidx is None:
        return
    header = [str(h) if h not in (None, "") else "" for h in rows[hidx]]
    findings = []

    for i, row in enumerate(rows):
        if i <= hidx:
            continue
        for col_idx, value in enumerate(row):
            if value in (None, ""):
                continue
            name = header[col_idx] if col_idx < len(header) else f"col{col_idx + 1}"
            line = i + 1
            if isinstance(value, str):
                if value != value.strip():
                    findings.append((line, name, "leading/trailing whitespace"))
                if re.search(r"password|passwd|pwd\s*[:=]", value, re.I):
                    findings.append((line, name, "possible credential in cell (value not reproduced)"))
                if re.search(r"my\.sharepoint\.com/(:[a-z]:/[a-z]/)?personal/", value, re.I):
                    findings.append((line, name, "personal OneDrive link"))
                if re.search(r"\bcademian|\bAcademain\b", value):
                    findings.append((line, name, "possible name typo (cademian/Academain)"))
                if re.search(r"\b\d{2,3}\+?\s+(experts|FTE|employees|people)", value, re.I):
                    findings.append((line, name, "numeric company claim in text (stats check candidate)"))
                if re.search(r"https?://", value) and not re.search(r"\.(pdf|pptx?|docx?|mp4|xlsx?)\b", value, re.I):
                    findings.append((line, name, "link without an obvious file extension"))
            if is_serial_date(value):
                findings.append((line, name, f"date serial {value} -> {serial_to_iso(value)}"))
            if isinstance(value, (dt.datetime, dt.date)):
                findings.append((line, name, "native date value (no conversion needed)"))

    if findings:
        out.append(f"#### Findings in `{ws.title}` ({len(findings)})")
        out.append("")
        out.append("| Row | Column | Finding |")
        out.append("|---|---|---|")
        for line, name, msg in findings:
            out.append(f"| {line} | {name} | {msg} |")
        out.append("")


def dump_sheet(ws):
    print(f"\n===== {ws.title} [{ws.sheet_state}] dims={ws.dimensions} =====")
    for i, row in enumerate(sheet_rows(ws), start=1):
        if any(c not in (None, "") for c in row):
            cells = [f"{c!r}" for c in row]
            while cells and cells[-1] == "None":
                cells.pop()
            print(f"R{i}: " + " | ".join(cells))


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("path")
    parser.add_argument("--dump", action="store_true", help="print raw non-empty rows")
    args = parser.parse_args()

    wb = load_workbook(args.path, data_only=True)
    if args.dump:
        for ws in wb.worksheets:
            dump_sheet(ws)
        return 0

    out = [f"# Workbook profile: `{args.path}`", ""]
    out.append("Sheets (in workbook order):")
    out.append("")
    for ws in wb.worksheets:
        out.append(f"- `{ws.title}` — {ws.sheet_state}, {ws.dimensions}")
    out.append("")
    for ws in wb.worksheets:
        profile_sheet(ws, out)
        find_issues(ws, out)
    sys.stdout.write("\n".join(out) + "\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

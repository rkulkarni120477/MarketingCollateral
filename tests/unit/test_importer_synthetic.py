"""Importer behaviour on a small synthetic workbook.

The workbook is built in a temp folder with invented values. A fake credential is
planted in Home!B18 so the tests can prove it never reaches any output.
"""
import json

import openpyxl
import pytest
from openpyxl.comments import Comment

from migration.importWorkbook.importer import run_dry_run
from migration.importWorkbook.reports import markdown_report, sanitized_fixture

FAKE_SECRET = "TEST-ONLY-SECRET-9f2c"


def _build(path, fte=(70, 85, 50, 75, 60, 45, 225, 90), extra_rows=True, fte_extra=None):
    wb = openpyxl.Workbook()
    home = wb.active
    home.title = "Home"
    home["B18"] = f"password: {FAKE_SECRET}"

    logos = wb.create_sheet("Logos_Stats")
    logos["B1"] = "Total Logos"
    logos["D1"] = "Last updated on: 24 Aug 2026"
    logos["F2"] = "FTEs"
    logos["G2"] = "=SUM(G3:G10)"
    logos["H2"] = "Approver Example"
    for i, value in enumerate(fte):
        logos[f"F{3 + i}"] = f"Role {i}"
        logos[f"G{3 + i}"] = value
    logos["F11"] = "Outside Role"
    logos["G11"] = 50
    logos["H11"] = "TBD"
    logos["F13"] = "Flex"
    logos["G13"] = "500+"
    logos["I13"] = "2026-09-01"
    logos["B3"] = "Client A"
    logos["C3"] = "Client B"
    logos["D3"] = "Client A"
    if fte_extra is not None:
        logos["G12"] = fte_extra

    pres = wb.create_sheet("01-PresentationsCatalogue")
    headers = ["Presentation Title", "Current Dir", "New Name", "Domain", "Short Description", "Puneet",
               "Review comments", "Slide Count/Duration", "Version", "Status", "Source (PPT)", "Published (PDF)",
               "Last Updated", "Last Updated By"]
    for col, text in enumerate(headers, start=1):
        pres.cell(row=1, column=col, value=text)
    pres["A2"] = "Deck One.pptx​"
    pres["A2"].hyperlink = "https://tenant.example/personal/deck_one.pptx"
    pres["C2"] = "Deck One"
    pres["D2"] = "HE+WD"
    pres["J2"] = None
    pres["B2"] = "Folder A"
    pres["B2"].hyperlink = "Folder A/relative"
    pres["M2"] = 46000
    pres["A3"] = "Deck Two.pdf"
    pres["D3"] = "AD"
    pres["J3"] = "Published"
    pres["A3"].hyperlink = "https://tenant.example/personal/deck_two.pdf"
    pres["A3"].comment = Comment("internal note for test only", "tester")

    one = wb.create_sheet("02-OnePager-Catalogue")
    one["A1"], one["B1"], one["D1"], one["E1"], one["F1"] = "One Pager", "Domain", "Status", "Source (PPT)", "Published (PDF)"
    one["A2"] = "One A"
    one["B2"] = "K12"
    one["D2"] = "Published"
    one["E2"] = "one_a.pptx"
    one["E2"].hyperlink = "https://tenant.example/personal/one_a.pptx"

    brochure = wb.create_sheet("03-Brochure-Catalogue")
    brochure["A1"], brochure["C1"], brochure["F1"] = "Brochure Name", "Domain", "Status"
    for name in ("04-CaseStudiesCatalogue", "05-MarketingVideos", "06-RateCards"):
        sheet = wb.create_sheet(name)
        sheet["A1"] = "Title"

    hidden = wb.create_sheet("Sheet1")
    hidden.sheet_state = "hidden"
    hidden["B2"] = "Deck One"
    hidden["H2"] = "Legacy update mentions 550+ claims"
    hidden["B3"] = "Totally unknown legacy name"
    if extra_rows:
        hidden["G4"] = "Note only, no name"
    onehidden = wb.create_sheet("02-OnePager-Catalogue_001")
    onehidden.sheet_state = "hidden"
    onehidden["A2"] = "Legacy One"

    wb.save(path)


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    _build(path)
    return path


def test_source_is_unchanged_and_secret_never_reaches_output(workbook, tmp_path):
    model = run_dry_run(workbook)
    assert model["workbook"]["sourceUnchanged"] is True
    assert model["exclusions"] == [{"sheet": "Home", "cell": "B18", "reason": "plaintext credential"}]
    serialized = json.dumps(model, default=str)
    assert FAKE_SECRET not in serialized
    assert FAKE_SECRET not in markdown_report(model)
    assert FAKE_SECRET not in json.dumps(sanitized_fixture(model), default=str)


def test_counts_status_and_domain_normalisation(workbook):
    model = run_dry_run(workbook)
    presentations = model["catalogue"]["Presentations"]
    assert len(presentations) == 2
    statuses = {r["title"]: r["status"] for r in presentations}
    assert statuses["Deck One.pptx"] == "Not specified"
    assert statuses["Deck Two.pdf"] == "Published"
    deck_one = next(r for r in presentations if r["title"].startswith("Deck One"))
    assert deck_one["domains"] == ["Higher Education", "Workforce Development"]
    assert "NeedsStatusReview" in deck_one["flags"]
    assert deck_one["titleLink"]["kind"] == "absolute"
    assert deck_one["currentDirLink"]["kind"] == "relative-unresolved"
    assert deck_one["fileTypeHint"] == "pptx"


def test_hidden_rows_do_not_become_catalogue_records(workbook):
    model = run_dry_run(workbook)
    assert model["totalVisibleRecords"] == 3
    assert model["hiddenEvidence"]["notesOnlyRows"] == 1
    assert model["hiddenEvidence"]["results"]["Sheet1"]["candidate-unique"] == 1
    assert model["hiddenEvidence"]["results"]["Sheet1"]["unmatched"] == 1


def test_comments_are_presence_only(workbook):
    model = run_dry_run(workbook)
    deck_two = next(r for r in model["catalogue"]["Presentations"] if r["title"] == "Deck Two.pdf")
    assert "internal note" not in json.dumps(model, default=str)
    assert deck_two["fileTypeHint"] == "pdf"


def test_fte_total_and_percentages_recalculate(tmp_path):
    path = tmp_path / "a.xlsx"
    _build(path, fte=(10, 10, 20))
    model = run_dry_run(path)
    stats = model["statistics"]
    assert stats["fteTotal"] == 40
    shares = [c["percentOfTotal"] for c in stats["fteComponents"] if c["numeric"]]
    assert shares[2] == pytest.approx(0.5)
    assert sum(shares) == pytest.approx(1.0)


def test_excluded_rows_do_not_enter_total(tmp_path):
    path = tmp_path / "b.xlsx"
    _build(path, fte=(10, 10, 20), fte_extra=999)  # G12 sits outside SUM(G3:G10)
    stats = run_dry_run(path)["statistics"]
    assert stats["fteTotal"] == 40
    assert all(item["inTotal"] is False for item in stats["outsideTotal"])


def test_zero_total_does_not_divide(tmp_path):
    path = tmp_path / "c.xlsx"
    _build(path, fte=(0, 0))
    stats = run_dry_run(path)["statistics"]
    assert stats["fteTotal"] == 0
    assert all(c["percentOfTotal"] is None for c in stats["fteComponents"])


def test_header_date_mismatch_is_flagged(workbook):
    stats = run_dry_run(workbook)["statistics"]
    assert stats["headerDate"] == "2026-08-24"
    assert "HeaderDateMismatchStatisticDate" in stats["flags"]
    flex = next(q for q in stats["qualifiedStatistics"] if q["label"] == "Flex")
    assert flex["value"]["qualifier"] == "at least" and flex["approval"] == "Unknown"


def test_client_counts_are_per_domain_and_distinct(workbook):
    clients = run_dry_run(workbook)["clients"]
    # B3 "Client A" (K12), C3 "Client B" (HE), D3 "Client A" (WD)
    assert clients["countsByDomain"] == {"K12": 1, "HE": 1, "WD": 1}
    assert clients["totalEntries"] == 3
    assert clients["distinctClients"] == 2
    assert clients["multiDomainClients"] == ["clienta"]
